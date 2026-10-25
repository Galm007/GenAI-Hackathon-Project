import re
from pathlib import Path
from typing import get_args

from etl import load_cds
from etl.load_cds import (
    C7Factors,
    C7Level,
    CDSCounts,
    CDSExtract,
    CDSTests,
    EarlyPlans,
    GPABucket,
    TransferProfile,
    extract_cds,
    sanity_check,
    section_texts,
    to_update,
)

MIGRATION = Path(__file__).resolve().parents[3] / "supabase" / "migrations" / "0001_extensions_types.sql"


def make_pdf(pages: list[list[tuple[int, int, str]]]) -> bytes:
    """Minimal PDF with a text layer: each page is a list of (x, y, text) runs in Helvetica 10."""
    objs: list[bytes] = []
    n_pages = len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n_pages))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    for i, runs in enumerate(pages):
        content = "BT /F1 10 Tf " + " ".join(f"1 0 0 1 {x} {y} Tm ({t}) Tj" for x, y, t in runs) + " ET"
        font = "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
                    f"/Resources << /Font << /F1 {font} >> >> >>".encode())
        objs.append(f"<< /Length {len(content)} >>\nstream\n{content}\nendstream".encode())
    out = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


CDS_PDF = make_pdf([
    [(72, 700, "General information and cover page")],
    [(72, 700, "C1 Applications"), (72, 680, "Total first-time applicants 50000")],
    [(72, 700, "Spillover page for C1")],
    [(72, 700, "C7 Relative importance"), (72, 680, "Factor"), (250, 680, "Very Important"), (400, 680, "Important"),
     (72, 660, "Rigor of secondary school record"), (260, 660, "X"), (72, 640, "Essay"), (410, 640, "X")],
    [(72, 700, "C9 SAT percentiles"), (72, 680, "C12 Average GPA 3.9")],
    [(72, 700, "Footer page")],
])


def test_section_texts_pick_marker_pages_and_next_page():
    t = section_texts(CDS_PDF)
    assert "Total first-time applicants 50000" in t["counts"] and "Spillover page" in t["counts"]
    assert "Rigor of secondary school record" in t["c7"] and "Average GPA" not in t["c7"].split("C9")[0]
    assert "C9 SAT percentiles" in t["tests"]


def test_layout_text_keeps_mark_columns_apart():
    c7 = section_texts(CDS_PDF)["c7"]
    rigor = next(line for line in c7.splitlines() if "Rigor" in line)
    essay = next(line for line in c7.splitlines() if "Essay" in line)
    assert rigor.index("X") < essay.index("X")             # "Very Important" column is left of "Important"


def mk_counts(**over):
    base = dict(cds_year="2024-2025", applicants_total=50000, admits_total=20000, enrolled_total=7000,
                ed_applicants=None, ed_admits=None, early_plans=EarlyPlans(ed=None, ed2=None, ea=True, rea=None),
                transfer=TransferProfile(accepts_transfers=True, applicants=None, admits=None, enrolled=None,
                                         min_college_gpa=None))
    return CDSCounts(**(base | over))


def mk_tests(**over):
    base = dict(test_policy="optional", sat_composite_25=1340, sat_composite_50=1430, sat_composite_75=1520,
                sat_ebrw_25=650, sat_ebrw_75=740, sat_math_25=690, sat_math_75=790, act_composite_25=30,
                act_composite_75=34, pct_submitting_sat=0.55, pct_submitting_act=0.30,
                gpa_distribution=[GPABucket(bucket="3.75-3.99", pct=0.5)], avg_hs_gpa=3.9, pct_top10_class=0.8)
    return CDSTests(**(base | over))


def mk_c7(**over):
    keys = list(C7Factors.model_fields)
    base = dict.fromkeys(keys)
    return C7Factors(**(base | {"rigor": "very_important", "essay": "important"} | over))


def extract(c=None, t=None, k=None):
    both = {**(c or mk_counts()).model_dump(), **(t or mk_tests()).model_dump()}
    return CDSExtract(**both, c7=k or mk_c7(), extraction_notes=[])


def test_sanity_check_keeps_good_data():
    x = sanity_check(extract())
    assert x.admits_total == 20000 and x.sat_composite_75 == 1520 and x.extraction_notes == []


def test_sanity_check_drops_impossible_values():
    x = sanity_check(extract(
        mk_counts(applicants_total=100, admits_total=900, enrolled_total=50),
        mk_tests(sat_composite_25=1500, sat_composite_75=1300, act_composite_75=40, pct_submitting_sat=63,
              pct_top10_class=-1, avg_hs_gpa=39.0, gpa_distribution=[GPABucket(bucket="4.0", pct=7000)]),
    ))
    assert x.applicants_total is None and x.admits_total is None
    assert x.sat_composite_25 is None and x.sat_composite_75 is None
    assert x.act_composite_25 is None and x.act_composite_75 is None
    assert x.pct_submitting_sat == 0.63 and x.pct_top10_class is None
    assert x.avg_hs_gpa is None and x.gpa_distribution == []
    assert len(x.extraction_notes) >= 5


def test_to_update_skips_nulls_and_keeps_c7():
    upd = to_update(sanity_check(extract(t=mk_tests(sat_composite_50=None))), "http://cds.pdf")
    assert upd["admit_rate"] == 0.4 and upd["c7_weights"] == {"rigor": "very_important", "essay": "important"}
    assert "sat_composite_50" not in upd and upd["source_urls"] == {"cds": "http://cds.pdf"}
    assert upd["data_quality"]["extractor"] == load_cds.settings.cds_model


async def test_extract_cds_makes_three_calls_and_assembles(monkeypatch):
    seen = []

    async def fake_chat(model_cls, system, user, **kw):
        seen.append((model_cls, kw["model"]))
        return {CDSCounts: mk_counts(), CDSTests: mk_tests(), C7Factors: mk_c7()}[model_cls]

    monkeypatch.setattr(load_cds, "chat_structured", fake_chat)
    x = await extract_cds(CDS_PDF, "uiuc")
    assert [m for m, _ in seen] == [CDSCounts, CDSTests, C7Factors]
    assert x is not None and x.c7.rigor == "very_important" and x.admits_total == 20000


async def test_pdf_without_text_layer_is_skipped(monkeypatch):
    async def boom(*a, **k):
        raise AssertionError("model must not be called")

    monkeypatch.setattr(load_cds, "chat_structured", boom)
    assert await extract_cds(make_pdf([[]]), "uiuc") is None


def test_c7_levels_match_the_database_enum():
    enum = re.search(r"c7_level\s+as enum \(([^)]*)\)", MIGRATION.read_text()).group(1)
    assert set(get_args(C7Level)) == set(re.findall(r"'(\w+)'", enum))


def make_xlsx(c7_levels: list[str]) -> bytes:
    """A miniature of the official CDS Excel template: item code | question | answer."""
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    c = wb.active
    c.title = "CDS-C"
    c.append(["C1: First-time applicants", None, None])
    c.append(["C117", "Total first-time, first-year students who applied", 73742.0])
    c.append(["C118", "Total first-time, first-year students who were admitted", 31247])
    c.append(["C102", "Total first-time, first-year women who applied", 5])           # per-gender rows are skipped
    c.append(["C905", "SAT Composite: 25th Percentile", 1390])
    for i, level in enumerate(c7_levels, 1):
        c.append([f"C7{i:02d}", f"factor {i}", level])
    c.append(["C2106", "Number of early decision applications", None])               # unanswered: dropped
    wb.create_sheet("CDS Definitions").append(["C117", "definition text that must be ignored", "x"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


LEVELS = ["Very Important"] * 3 + ["Important"] * 5 + ["Considered"] * 5 + ["Not Considered"] * 5


def test_xlsx_is_detected_and_sections_keep_only_answered_rows():
    data = make_xlsx(LEVELS)
    assert load_cds.is_xlsx(data) and not load_cds.is_xlsx(CDS_PDF)
    t = section_texts(data)
    assert "C117 | Total first-time, first-year students who applied | 73742" in t["counts"]
    assert "C102" not in t["counts"] and "C2106" not in t["counts"] and "definition" not in t["counts"]
    assert "C905 | SAT Composite: 25th Percentile | 1390" in t["tests"]


def test_xlsx_c7_is_read_by_item_code_in_field_order():
    c7 = load_cds.xlsx_c7(make_xlsx(LEVELS))
    assert c7.rigor == "very_important" and c7.academic_gpa == "very_important"
    assert c7.test_scores == "important" and c7.applicant_interest == "not_considered"
    assert load_cds.xlsx_c7(make_xlsx(["Not a level"] * 18)).rigor is None


async def test_xlsx_extraction_skips_the_c7_model_call(monkeypatch):
    seen = []

    async def fake_chat(model_cls, system, user, **kw):
        seen.append(model_cls)
        return {CDSCounts: mk_counts(), CDSTests: mk_tests()}[model_cls]

    monkeypatch.setattr(load_cds, "chat_structured", fake_chat)
    x = await extract_cds(make_xlsx(LEVELS), "uiuc")
    assert seen == [CDSCounts, CDSTests] and x is not None and x.c7.rigor == "very_important"


def test_google_sheet_links_become_xlsx_exports():
    sheet = "https://docs.google.com/spreadsheets/d/AbC-123_x/edit?usp=sharing"
    assert load_cds.download_url(sheet) == "https://docs.google.com/spreadsheets/d/AbC-123_x/export?format=xlsx"
    assert load_cds.download_url("https://example.edu/cds.pdf") == "https://example.edu/cds.pdf"


def test_xlsx_test_policy_is_read_from_c802():
    import io

    import openpyxl

    def book(answer):
        wb = openpyxl.Workbook()
        wb.active.append(["C802", "SAT or ACT", answer])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    assert load_cds.xlsx_test_policy(book("Not considered for admission, even if submitted")) == "blind"
    assert load_cds.xlsx_test_policy(book("Not required for admission, but considered if submitted")) == "optional"
    assert load_cds.xlsx_test_policy(book("Required for some")) == "required"
    assert load_cds.xlsx_test_policy(book(None)) is None


HEADER = "Academic" + " " * 17 + "Very important    Important    Considered    Not considered"
COL = {"very_important": 32, "important": 47, "considered": 60, "not_considered": 75}     # header centres


def grid_row(label: str, level: str) -> str:
    return label + " " * (COL[level] - len(label)) + "X"


def test_pdf_c7_reads_the_column_of_each_x():
    rows = [HEADER, grid_row("Rigor of secondary school record"[:30], "very_important"),
            grid_row("Class rank", "not_considered"), grid_row("Academic GPA", "important"),
            grid_row("Standardized test scores", "considered")]
    rows += [grid_row(lb, "very_important") for lb in (
        "Application essay", "Recommendation(s)", "Interview", "Extracurricular activities", "Talent/ability",
        "Character/personal qualities", "First generation", "Alumni/ae relation", "Geographical residence",
        "State residency")]
    rows.append(grid_row("Volunteer work", "important"))
    c7 = load_cds.pdf_c7("\n".join(rows))
    assert c7 is not None
    assert (c7.rigor, c7.class_rank, c7.academic_gpa, c7.test_scores) == (
        "very_important", "not_considered", "important", "considered")
    assert c7.essay == "very_important" and c7.volunteer_work == "important"


def test_pdf_c7_reads_form_field_levels_at_the_end_of_the_row():
    labels = ["Rigor of secondary school record", "Class rank", "Academic Grade Point Average (GPA)",
              "Recommendations", "Standardized test scores", "Application essay", "Interview",
              "Extracurriculuar activities", "Talent/ability", "Character/personal qualities",
              "First generation", "Alumni/ae relation", "Geographical residence", "State residency",
              "Religious affilitation/commitment", "Volunteer work", "Work experience"]
    text = "\n".join(f"{lb}                    {'Not Considered' if i % 2 else 'Very Important'}"
                     for i, lb in enumerate(labels))
    c7 = load_cds.pdf_c7(text)
    assert c7 is not None and c7.rigor == "very_important" and c7.class_rank == "not_considered"
    assert c7.extracurriculars == "not_considered" and c7.applicant_interest is None


def test_pdf_c7_gives_up_on_an_unfamiliar_layout():
    assert load_cds.pdf_c7("C7 Relative importance ...\nRigor of secondary school record X") is None


def test_pdf_c1_totals_take_the_unsplit_row_and_its_largest_number():
    text = """\
Total first-time, first-year men who applied                    40,022
Total first-time, first-year women who applied                  19,765
Total first-time, first-year who applied      59,789   10,674   38,320   10,795   0
Total first-time, first-year who were admitted     8,413   3,536   3,992   885   0
Total first-time, first-year who enrolled      3,850   2,353   1,164   333   0
"""
    assert load_cds.pdf_c1_totals(text) == (59789, 8413, 3850)
    impossible = "\n".join(f"Total first-time, first-year who {k}   {n}" for k, n in
                           (("applied", 10), ("were admitted", 20), ("enrolled", 5)))
    assert load_cds.pdf_c1_totals(impossible) is None                  # admits > applicants: leave it to the model


def test_pdf_percentiles_read_the_first_three_number_row_per_test():
    text = """\
              25th Percentile   50th Percentile   75th Percentile
SAT Composite (400 - 1600)        1360    1460    1530
SAT Math                          680     740     780
ACT Composite                     31      33      34
ACT Writing                       -       -       -
              SAT Composite
SAT Composite                     9999    9999    9999
"""
    p = load_cds.pdf_percentiles(text)
    assert (p["sat_composite_25"], p["sat_composite_75"], p["sat_math_50"], p["act_composite_75"]) == (
        1360, 1530, 740, 34)
    assert "act_writing_25" not in p


def test_xlsx_tests_read_by_label_and_treat_zero_as_not_reported():
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["C905", "SAT Composite: 25th Percentile", 1390])
    ws.append(["C914", "ACT Composite: 25th Percentile", 0])
    ws.append(["C901", "Submitting SAT Scores", 0.4128])
    ws.append(["C1001", "Percent in top tenth of high school graduating class", 56.07])
    ws.append(["C1201", "Average high school GPA of all degree-seeking, first-time", 0])
    buf = io.BytesIO()
    wb.save(buf)
    t = load_cds.xlsx_tests(buf.getvalue())
    assert t["sat_composite_25"] == 1390 and t["act_composite_25"] is None and t["avg_hs_gpa"] is None
    assert t["pct_submitting_sat"] == 0.4128 and t["pct_top10_class"] == 56.07
